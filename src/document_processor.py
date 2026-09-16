import os
import PyPDF2
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter

class DocumentProcessor:
    def __init__(self, chunk_size=600, chunk_overlap=100):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            is_separator_regex=False,
        )

    def extract_text_from_pdf(self, file_path):
        text = ""
        page_count = 0
        try:
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                page_count = len(reader.pages)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            print(f"Error reading PDF {file_path}: {e}")
        return text, page_count

    def extract_text_from_docx(self, file_path):
        text = ""
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        except Exception as e:
            print(f"Error reading Word document {file_path}: {e}")
        # Word documents don't have a reliable page count in docx library without rendering.
        # We'll approximate or just return 0.
        return text, 0

    def process_document(self, file_path):
        """Reads a document and returns a tuple (chunks, raw_text, page_count)."""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.pdf':
            text, page_count = self.extract_text_from_pdf(file_path)
        elif ext in ['.doc', '.docx']:
            text, page_count = self.extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")
        
        if not text.strip():
            return [], "", 0
            
        chunks = self.text_splitter.split_text(text)
        return chunks, text, page_count

    def process_uploaded_file(self, uploaded_file):
        """Reads a streamlit uploaded file object and returns a tuple (chunks, raw_text, page_count)."""
        ext = os.path.splitext(uploaded_file.name)[1].lower()
        text = ""
        page_count = 0
        if ext == '.pdf':
            try:
                reader = PyPDF2.PdfReader(uploaded_file)
                page_count = len(reader.pages)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            except Exception as e:
                print(f"Error reading uploaded PDF {uploaded_file.name}: {e}")
        elif ext in ['.doc', '.docx']:
            try:
                doc = docx.Document(uploaded_file)
                for para in doc.paragraphs:
                    text += para.text + "\n"
                # Approximation for docx pages: 1 page per 3000 chars roughly.
                page_count = max(1, len(text) // 3000)
            except Exception as e:
                print(f"Error reading uploaded Word document {uploaded_file.name}: {e}")
        else:
            raise ValueError(f"Unsupported file format: {ext}")
        
        if not text.strip():
            return [], "", 0
            
        chunks = self.text_splitter.split_text(text)
        return chunks, text, page_count
